"""HTTP/storage integration tests. Stub inference is NOT real-model validation."""

import copy
import io
import json
import sqlite3
import subprocess
import sys
import threading
import time
import uuid
import wave
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.app import main, ml, storage


def poll(client, meeting_id):
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        response = client.get(f"/api/meetings/{meeting_id}")
        assert response.status_code == 200
        meeting = response.json()
        if meeting["status"] in ("review_ready", "failed"):
            return meeting
        time.sleep(0.01)
    pytest.fail("Processing did not finish")


def demo(client):
    response = client.post("/api/meetings/demo")
    assert response.status_code == 200, response.text
    return response.json()


def reviewed(client, meeting):
    actions = copy.deepcopy(meeting["actions"])
    for action in actions:
        action.update(assignee="Әлия", review_reasons=[])
    response = client.patch(f"/api/meetings/{meeting['id']}", json={"revision": meeting["revision"], "actions": actions})
    assert response.status_code == 200, response.text
    return response.json()


def approved(client, meeting):
    response = client.post(f"/api/meetings/{meeting['id']}/approve", json={"revision": meeting["revision"]})
    assert response.status_code == 200, response.text
    return response.json()


def ml_result():
    return {"segments": [
        {"id": "segment", "start_ms": 0, "end_ms": 0, "speaker_id": "", "text": "Әлия дайындасын. Срок не указан."},
        {"id": "extra", "start_ms": 0, "end_ms": 0, "speaker_id": "SPEAKER_01", "text": "Қазақша: ә ғ қ ң ө ұ ү һ і. Дополнительный контекст."}],
        "summary": "Подготовить отчёт. Әлия тексереді.",
        "actions": [{"id": "deterministic-ml-id", "title": "Подготовить отчёт", "assignee": "Әлия", "due_text": "Срок не указан",
                     "due_date": None, "evidence_segment_ids": ["segment"], "review_reasons": [], "status": "open"}],
        "warnings": []}


def text_payload():
    return {"title": "Проверка", "occurred_at": "2026-09-20", "timezone": "Asia/Almaty", "participants": [],
            "transcript": "Әлия дайындасын. Срок не указан."}


def wav_bytes():
    output = io.BytesIO()
    with wave.open(output, "wb") as wav:
        wav.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
        wav.writeframes(b"\0\0" * 1600)
    return output.getvalue()


def upload(client, content=None, filename="meeting.wav", **extra):
    return client.post("/api/meetings/audio", data={"title": "Запись", "occurred_at": "2026-09-20", "timezone": "Asia/Almaty", "participants": "[]", **extra},
                       files={"file": (filename, wav_bytes() if content is None else content, "audio/wav")})


def test_review_approval_versions_and_immutable_export(client):
    meeting = demo(client)
    url = f"/api/meetings/{meeting['id']}"
    assert client.get(url + "/export.docx").status_code == 409
    assert client.post(url + "/approve", json={"revision": meeting["revision"]}).status_code == 409
    meeting = reviewed(client, meeting)
    previous = meeting["revision"]
    meeting = approved(client, meeting)
    assert client.patch(url, json={"revision": previous, "summary": "stale"}).status_code == 409
    first = client.get(url + "/export.docx")
    assert first.status_code == 200
    first_xml = Document(io.BytesIO(first.content))._element.xml
    assert "ДЕМОНСТРАЦИОННЫЕ" in first_xml
    assert f"Утверждённая редакция: {meeting['revision']}" in first_xml
    action_id = meeting["actions"][0]["id"]
    response = client.patch(f"/api/actions/{action_id}", json={"status": "done"})
    assert response.status_code == 200
    assert response.json()["id"] == action_id
    latest = client.get(url).json()
    assert latest["approved"] and latest["revision"] == meeting["revision"] + 1
    assert client.patch(url, json={"revision": meeting["revision"], "actions": meeting["actions"]}).status_code == 409
    assert Document(io.BytesIO(client.get(url + "/export.docx").content))._element.xml == first_xml
    assert client.patch(f"/api/actions/{action_id}", json={"status": "done"}).status_code == 200
    assert client.get(url).json()["revision"] == latest["revision"]
    changed = client.patch(url, json={"revision": latest["revision"], "summary": "Новое содержание"}).json()
    assert not changed["approved"]
    assert client.get(url + "/export.docx").status_code == 409
    assert client.patch(f"/api/actions/{action_id}", json={"status": "open"}).status_code == 409
    changed = approved(client, changed)
    xml = Document(io.BytesIO(client.get(url + "/export.docx").content))._element.xml
    assert "Новое содержание" in xml
    assert f"Утверждённая редакция: {changed['revision']}" in xml


def test_parallel_stale_edits_are_atomic(client):
    meeting = demo(client)
    barrier = threading.Barrier(2)

    def write(summary):
        barrier.wait()
        return client.patch(f"/api/meetings/{meeting['id']}", json={"revision": meeting["revision"], "summary": summary}).status_code

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(write, ("First", "Second")))
    assert sorted(results) == [200, 409]
    assert client.get(f"/api/meetings/{meeting['id']}").json()["revision"] == meeting["revision"] + 1


@pytest.mark.parametrize("change", [
    {"due_date": "2026-02-30"}, {"due_date": "20260923"}, {"status": "unknown"},
    {"evidence_segment_ids": ["missing"]}, {"evidence_segment_ids": [42]},
    {"title": "control\u0000character"}, {"review_reasons": [""]}, {"assignee": None},
])
def test_invalid_edits_do_not_change_data(client, change):
    meeting = demo(client)
    actions = copy.deepcopy(meeting["actions"])
    actions[0].update(change)
    url = f"/api/meetings/{meeting['id']}"
    response = client.patch(url, json={"revision": meeting["revision"], "summary": "must roll back", "actions": actions})
    assert response.status_code == 422
    assert isinstance(response.json()["detail"], str)
    assert client.get(url).json() == meeting


def test_unique_ids_and_edit_identity(client):
    first, second = demo(client), demo(client)
    assert not {a["id"] for a in first["actions"]} & {a["id"] for a in second["actions"]}
    actions = copy.deepcopy(second["actions"])
    actions[0]["id"] = first["actions"][0]["id"]
    url = f"/api/meetings/{second['id']}"
    assert client.patch(url, json={"revision": second["revision"], "actions": actions}).status_code == 409
    assert client.get(url).json() == second
    duplicate = [second["actions"][0], second["actions"][0]]
    assert client.patch(url, json={"revision": second["revision"], "actions": duplicate}).status_code == 422
    edited = reviewed(client, second)
    assert [a["id"] for a in edited["actions"]] == [a["id"] for a in second["actions"]]


def test_audio_upload_stages_review_docx_and_playback(client, monkeypatch):
    observed = []

    def process(path, context, stage):
        assert Path(path).read_bytes() == wav_bytes()
        assert Path(path).parent == storage.UPLOAD_DIR
        assert context == {"title": "Запись", "occurred_at": "2026-09-20", "timezone": "Asia/Almaty", "participants": []}
        for name in ("transcribing", "diarizing", "extracting"):
            stage(name)
            observed.append(storage.list_meetings()[0]["status"])
        result = ml_result()
        result["segments"][0]["end_ms"] = 100
        return result

    monkeypatch.setattr(ml, "process_audio", process)
    response = upload(client, filename="../../original.WAV")
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "queued"
    meeting = poll(client, response.json()["id"])
    assert meeting["status"] == "review_ready", meeting["error"]
    assert observed == ["transcribing", "diarizing", "extracting"]
    assert meeting["occurred_at"] == "2026-09-20" and meeting["has_audio"]
    url = f"/api/meetings/{meeting['id']}"
    audio = client.get(url + "/audio", headers={"Range": "bytes=0-43"})
    assert audio.status_code == 206
    assert audio.content == wav_bytes()[:44]
    meeting = approved(client, meeting)
    response = client.get(url + "/export.docx")
    document = Document(io.BytesIO(response.content))
    xml = document._element.xml
    assert len(document.tables[0].rows) == 2
    assert "ә ғ қ ң ө ұ ү һ і" in xml
    assert "Дополнительный контекст" in xml
    assert "Дата не указана" in xml
    assert "ДЕМОНСТРАЦИОННЫЕ" not in xml


def test_text_without_speaker_and_repeated_ml_ids(client, monkeypatch):
    monkeypatch.setattr(ml, "process_text", lambda *args: ml_result())
    meetings = []
    for _ in range(2):
        response = client.post("/api/meetings/text", json=text_payload())
        meetings.append(poll(client, response.json()["id"]))
    assert all(m["status"] == "review_ready" for m in meetings)
    assert meetings[0]["actions"][0]["id"] != meetings[1]["actions"][0]["id"]
    assert meetings[0]["segments"][0]["speaker_id"] == ""
    for meeting in meetings:
        uuid.UUID(meeting["actions"][0]["id"])
        assert client.get(f"/api/meetings/{meeting['id']}/audio").status_code == 404


@pytest.mark.parametrize("broken", ["exception", "evidence", "timestamps", "empty", "warnings"])
def test_ml_failure_is_explicit_and_never_demo(client, monkeypatch, broken):
    def process(*args):
        if broken == "exception":
            raise RuntimeError("Install local model weights")
        result = ml_result()
        if broken == "evidence":
            result["actions"][0]["evidence_segment_ids"] = ["invented"]
        elif broken == "timestamps":
            result["segments"][0]["end_ms"] = 100
        elif broken == "empty":
            result["segments"], result["actions"] = [], []
        else:
            result["warnings"] = ["bad\u0000"]
        return result

    monkeypatch.setattr(ml, "process_text", process)
    response = client.post("/api/meetings/text", json=text_payload())
    meeting = poll(client, response.json()["id"])
    assert meeting["status"] == "failed" and not meeting["approved"]
    assert meeting["source_mode"] == "text" and not meeting["actions"]
    assert meeting["error"].startswith("extracting:")
    assert meeting["revision"] == 1


def test_upload_validation_and_cleanup(client, monkeypatch):
    monkeypatch.setattr(main, "MAX_UPLOAD", 64)
    assert upload(client, b"x" * 65).status_code == 413
    assert upload(client, b"").status_code == 422
    assert upload(client, b"x", filename="wrong.exe").status_code == 422
    assert upload(client, b"x", participants="not-json").status_code == 422
    assert upload(client, b"x", timezone="unknown/zone").status_code == 422
    assert not list(storage.UPLOAD_DIR.iterdir())
    assert client.get("/api/meetings").json() == []


def test_request_errors_have_string_details(client):
    for response in (client.post("/api/meetings/text", content="bad", headers={"Content-Type": "application/json"}),
                     client.post("/api/meetings/audio"), client.get("/api/unknown")):
        assert response.status_code in (404, 422)
        assert isinstance(response.json()["detail"], str)


def test_notifications_use_local_dates_confirmed_tasks_and_stable_ids(client, monkeypatch):
    class Clock(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2026, 9, 23, 20, 30, tzinfo=timezone.utc).astimezone(tz)

    monkeypatch.setattr(main, "datetime", Clock)
    meeting = demo(client)
    base = meeting["actions"][0]
    due_dates = ["2026-09-23", "2026-09-24", "2026-09-27", "2026-09-28", None]
    actions = [{**base, "id": str(uuid.uuid4()), "assignee": "Әлия", "review_reasons": [], "due_date": due,
                "due_text": due or "после согласования", "status": "open"} for due in due_dates]
    url = f"/api/meetings/{meeting['id']}"
    meeting = client.patch(url, json={"revision": meeting["revision"], "actions": actions}).json()
    assert client.get("/api/notifications").json() == []
    meeting = approved(client, meeting)
    notices = client.get("/api/notifications").json()
    assert [n["kind"] for n in notices] == ["overdue", "due_soon", "due_soon"]
    assert notices == client.get("/api/notifications").json()
    assert len({n["id"] for n in notices}) == 3
    client.patch(f"/api/actions/{actions[0]['id']}", json={"status": "done"})
    assert len(client.get("/api/notifications").json()) == 2
    assert not next(a for a in client.get("/api/actions").json() if a["id"] == actions[0]["id"])["overdue"]
    meeting = client.get(url).json()
    client.patch(url, json={"revision": meeting["revision"], "summary": "Needs reapproval"})
    assert client.get("/api/notifications").json() == []


def test_due_date_in_export_is_authoritative(client):
    meeting = reviewed(client, demo(client))
    meeting["actions"][0].update(due_date="2026-12-01", due_text="после согласования")
    url = f"/api/meetings/{meeting['id']}"
    meeting = client.patch(url, json={"revision": meeting["revision"], "actions": meeting["actions"]}).json()
    approved(client, meeting)
    document = Document(io.BytesIO(client.get(url + "/export.docx").content))
    assert document.tables[0].rows[1].cells[3].text == "2026-12-01\nИсходный срок: после согласования"


def test_restart_persists_revision_status_snapshot_and_audio(isolated_storage, monkeypatch):
    monkeypatch.setattr(ml, "process_audio", lambda *args: ml_result())
    with TestClient(main.app) as client:
        response = upload(client)
        meeting = approved(client, poll(client, response.json()["id"]))
        url = f"/api/meetings/{meeting['id']}"
        doc = Document(io.BytesIO(client.get(url + "/export.docx").content))._element.xml
        client.patch(f"/api/actions/{meeting['actions'][0]['id']}", json={"status": "done"})
        latest = client.get(url).json()
    # A separate interpreter proves there is no dependence on process memory.
    code = "import sqlite3,json,sys; db=sqlite3.connect(sys.argv[1]); print(db.execute('SELECT data FROM meetings WHERE id=?',(sys.argv[2],)).fetchone()[0])"
    read = subprocess.run([sys.executable, "-c", code, str(storage.DB_PATH), meeting["id"]], capture_output=True, text=True, encoding="utf-8", check=True,
                          env={**__import__("os").environ, "PYTHONIOENCODING": "utf-8"})
    assert json.loads(read.stdout) == latest
    with TestClient(main.app) as client:
        assert client.get(url).json() == latest
        assert client.get(url + "/audio").content == wav_bytes()
        assert Document(io.BytesIO(client.get(url + "/export.docx").content))._element.xml == doc


def test_restart_fails_interrupted_jobs_and_resumes_only_queued(isolated_storage, monkeypatch):
    ids = {}
    for stage in ("queued", "transcribing", "diarizing", "extracting"):
        meeting = main._new_meeting("Restart", "2026-09-20", "Asia/Almaty", [], "text")
        meeting["status"] = stage
        storage.insert(meeting, transcript="Saved source")
        ids[stage] = meeting["id"]
    calls = []
    monkeypatch.setattr(ml, "process_text", lambda text, *args: (calls.append(text), ml_result())[1])
    with TestClient(main.app) as client:
        assert poll(client, ids["queued"])["status"] == "review_ready"
        for stage in ("transcribing", "diarizing", "extracting"):
            meeting = client.get(f"/api/meetings/{ids[stage]}").json()
            assert meeting["status"] == "failed"
            assert stage in meeting["error"] and "restart" in meeting["error"]
    assert calls == ["Saved source"]


def test_only_one_inference_runs_and_queued_source_is_durable(client, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    calls = []

    def process(text, *args):
        calls.append(text)
        entered.set()
        assert release.wait(5)
        return ml_result()

    monkeypatch.setattr(ml, "process_text", process)
    first = client.post("/api/meetings/text", json=text_payload()).json()
    try:
        assert entered.wait(2)
        second = client.post("/api/meetings/text", json={**text_payload(), "transcript": "Second source"}).json()
        assert len(calls) == 1
        assert storage.get_record(second["id"])["transcript"] == "Second source"
        assert client.get(f"/api/meetings/{second['id']}").json()["status"] == "queued"
    finally:
        release.set()
    assert poll(client, first["id"])["status"] == "review_ready"
    assert poll(client, second["id"])["status"] == "review_ready"
    assert len(calls) == 2


def test_legacy_database_upgrade_preserves_approved_protocol(isolated_storage):
    meeting = main._new_meeting("Legacy", "2026-09-20", "Asia/Almaty", [], "text")
    meeting.update(status="review_ready", approved=True, revision=4)
    with sqlite3.connect(storage.DB_PATH) as db:
        db.execute("DROP TABLE meetings")
        db.execute("CREATE TABLE meetings (id TEXT PRIMARY KEY, created_at TEXT NOT NULL, data TEXT NOT NULL, audio_path TEXT, transcript TEXT)")
        db.execute("INSERT INTO meetings VALUES (?, ?, ?, NULL, NULL)", (meeting["id"], meeting["created_at"], json.dumps(meeting)))
    storage.initialize()
    storage.initialize()
    assert storage.get_record(meeting["id"])["approved_meeting"] == meeting
