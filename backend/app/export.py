"""Render an approved meeting as a Word document."""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def render_docx(meeting: dict) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    title = doc.add_heading("Протокол встречи", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if meeting["source_mode"] == "demo":
        doc.add_paragraph("ДЕМОНСТРАЦИОННЫЕ ВЫМЫШЛЕННЫЕ ДАННЫЕ")
    doc.add_paragraph(f"Тема: {meeting['title']}")
    doc.add_paragraph(f"Дата: {meeting['occurred_at']} ({meeting['timezone']})")
    if meeting["source_mode"] == "text":
        doc.add_paragraph("Источник: импорт текста. Временные метки 00:00 не привязаны к аудио.")

    doc.add_heading("Участники", 1)
    for participant in meeting["participants"]:
        speaker = participant.get("speaker_id")
        doc.add_paragraph(participant["display_name"] + (f" — {speaker}" if speaker else ""), style="List Bullet")

    doc.add_heading("Краткое содержание", 1)
    doc.add_paragraph(meeting["summary"] or "—")

    doc.add_heading("Поручения", 1)
    if meeting["actions"]:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, heading in zip(table.rows[0].cells, ("№", "Поручение", "Ответственный", "Срок", "Статус")):
            cell.text = heading
        for i, action in enumerate(meeting["actions"], 1):
            row = table.add_row().cells
            for cell, value in zip(row, (str(i), action["title"], action["assignee"], action["due_text"] or (action["due_date"] or "Не указан"), "Выполнено" if action["status"] == "done" else "В работе")):
                cell.text = value
    else:
        doc.add_paragraph("Поручения не зафиксированы.")

    doc.add_heading("Основания в исходном материале", 1)
    evidence = {segment["id"]: segment for segment in meeting["segments"]}
    for i, action in enumerate(meeting["actions"], 1):
        if not action["evidence_segment_ids"]:
            continue
        doc.add_paragraph(f"Поручение {i}: {action['title']}")
        for segment_id in action["evidence_segment_ids"]:
            segment = evidence[segment_id]
            if meeting["source_mode"] == "audio":
                label = f"[{_clock(segment['start_ms'])}] {segment['speaker_id']}"
            else:
                label = f"{segment['speaker_id']} (без аудиометки)"
            doc.add_paragraph(f"{label}: {segment['text']}", style="List Bullet")

    if meeting["warnings"]:
        doc.add_heading("Примечания", 1)
        for warning in meeting["warnings"]:
            doc.add_paragraph(warning, style="List Bullet")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _clock(milliseconds: int) -> str:
    seconds = milliseconds // 1000
    return f"{seconds // 3600:02d}:{seconds // 60 % 60:02d}:{seconds % 60:02d}"
