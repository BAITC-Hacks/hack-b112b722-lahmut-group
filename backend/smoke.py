"""Manual real-model acceptance over HTTP; never substitutes demo or stub ML.

Upload a real recording, inspect/edit the saved JSON, then explicitly approve it.
Install backend/requirements-dev.txt for the httpx client.
"""

import argparse
import io
import json
import sys
import time
from pathlib import Path

import httpx
from docx import Document


def checked(response):
    if response.is_error:
        try:
            detail = response.json().get("detail", response.text)
        except ValueError:
            detail = response.text
        raise RuntimeError(f"HTTP {response.status_code}: {detail}")
    return response


def write_json(path, meeting):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(meeting, ensure_ascii=False, indent=2), encoding="utf-8")


def upload(client, args):
    health = checked(client.get("/api/health")).json()
    if not all(health["providers"].get(name) for name in ("speech", "diarization", "llm")):
        raise RuntimeError("Real audio providers are not ready: " + json.dumps(health["details"], ensure_ascii=False))
    participants = json.loads(args.participants.read_text(encoding="utf-8")) if args.participants else []
    with args.audio.open("rb") as source:
        meeting = checked(client.post("/api/meetings/audio", files={"file": (args.audio.name, source)}, data={
            "title": args.title, "occurred_at": args.occurred_at, "timezone": args.timezone,
            "participants": json.dumps(participants, ensure_ascii=False),
        })).json()
    write_json(args.output, meeting)
    print(f"Meeting {meeting['id']} saved. Review JSON: {args.output}")
    deadline, previous = time.monotonic() + args.timeout, None
    while True:
        meeting = checked(client.get(f"/api/meetings/{meeting['id']}")).json()
        if meeting["status"] != previous:
            print(meeting["status"], flush=True)
            previous = meeting["status"]
            write_json(args.output, meeting)
        if meeting["status"] == "failed":
            raise RuntimeError(meeting["error"])
        if meeting["status"] == "review_ready":
            break
        if time.monotonic() >= deadline:
            raise RuntimeError(f"Timed out waiting; meeting {meeting['id']} remains saved on the server. Inspect it through the UI/API.")
        time.sleep(1)
    if meeting["source_mode"] != "audio" or not meeting["segments"] or not meeting["has_audio"]:
        raise RuntimeError("Expected a real audio result with transcript segments")
    print("Inspect sources, assignees, dates and summary in the JSON. Correct them and clear only resolved review_reasons before the review command.")


def review(client, args):
    reviewed = json.loads(args.input.read_text(encoding="utf-8"))
    if reviewed["source_mode"] != "audio" or reviewed["status"] != "review_ready":
        raise RuntimeError("The real-audio check requires a review_ready audio meeting")
    url = f"/api/meetings/{reviewed['id']}"
    meeting = checked(client.patch(url, json={key: reviewed[key] for key in ("revision", "participants", "summary", "actions")})).json()
    write_json(args.input, meeting)
    if not args.approve:
        print("Corrections saved. Approval was not requested.")
        return
    meeting = checked(client.post(url + "/approve", json={"revision": meeting["revision"]})).json()
    write_json(args.input, meeting)
    response = checked(client.get(url + "/export.docx"))
    document = Document(io.BytesIO(response.content))
    if meeting["actions"] and (not document.tables or len(document.tables[0].rows) != len(meeting["actions"]) + 1):
        raise RuntimeError("DOCX is missing action rows")
    args.docx.parent.mkdir(parents=True, exist_ok=True)
    args.docx.write_bytes(response.content)
    print(f"Approved revision {meeting['revision']}; DOCX opens structurally: {args.docx}")
    print("Inspect the document visually. This check does not judge ASR/diarization/extraction quality.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    commands = parser.add_subparsers(dest="command", required=True)
    command = commands.add_parser("upload", help="Upload real audio, await inference, save review JSON")
    command.add_argument("--audio", type=Path, required=True)
    command.add_argument("--title", required=True)
    command.add_argument("--occurred-at", required=True, help="Date of the meeting, YYYY-MM-DD")
    command.add_argument("--timezone", default="Asia/Almaty")
    command.add_argument("--participants", type=Path, help="JSON array matching Participant in contracts/README.md")
    command.add_argument("--output", type=Path, default=Path("artifacts/real-review.json"))
    command.add_argument("--timeout", type=int, default=10800)
    command = commands.add_parser("review", help="Save manually reviewed JSON; optionally approve and export")
    command.add_argument("--input", type=Path, required=True)
    command.add_argument("--approve", action="store_true", help="Confirm you manually verified the content, evidence and unresolved questions")
    command.add_argument("--docx", type=Path, default=Path("artifacts/real-protocol.docx"))
    args = parser.parse_args()
    try:
        with httpx.Client(base_url=args.base_url.rstrip("/"), timeout=120, trust_env=False, follow_redirects=False) as client:
            (upload if args.command == "upload" else review)(client, args)
    except (RuntimeError, httpx.HTTPError, OSError, ValueError, KeyError) as exc:
        print(str(exc), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
